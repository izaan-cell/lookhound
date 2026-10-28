"""Checks an installed Lookhound end to end, inside the app's own Python:

    Lookhound.exe -c "exec(open(r'path/to/smoke_test.py').read())"

Set SMOKE_OUT (where results go), LOOKHOUND_HOME (a throwaway data folder) and
QT_QPA_PLATFORM=offscreen. It makes a few sample files, lets Lookhound index
them, runs searches, saves screenshots, and writes report.json with "passed".
"""
import json
import os
import time
import traceback
from pathlib import Path

OUT = Path(os.environ["SMOKE_OUT"])
HOME = Path(os.environ["LOOKHOUND_HOME"])
FILES = OUT / "files"
INDEX_TIMEOUT = 25 * 60  # the first run downloads about 1 GB of models

report = {"steps": [], "passed": False}


def step(name, ok, detail=""):
    report["steps"].append({"step": name, "ok": bool(ok), "detail": str(detail)})
    _save()


def _save():
    (OUT / "report.json").write_text(json.dumps(report, indent=2))


def make_samples():
    import io
    import tarfile
    import zipfile

    import py7zr

    from PIL import Image, ImageDraw
    from pptx import Presentation

    FILES.mkdir(parents=True, exist_ok=True)
    (FILES / "trip_notes.md").write_text(
        "# Lisbon trip\nFlights booked for June. Hotel near the river, tram 28 to the castle.")
    (FILES / "pasta_recipe.txt").write_text(
        "Tomato pasta: garlic, olive oil, basil, simmer the sauce for twenty minutes.")
    (FILES / "login.py").write_text(
        "def handle_login(request):\n    user = check_password(request.form['email'],"
        " request.form['password'])\n    return start_session(user)\n")
    image = Image.new("RGB", (256, 256), "white")
    ImageDraw.Draw(image).ellipse((40, 40, 216, 216), fill=(30, 80, 220))
    image.save(FILES / "blue_circle.png")
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Quarterly sales review"
    slide.placeholders[1].text = "Revenue by region, pricing and customer retention."
    deck.save(FILES / "sales_deck.pptx")
    with zipfile.ZipFile(FILES / "bali_2025.zip", "w") as zf:
        zf.writestr("beach_day/surfing_lesson.jpg", b"")
        zf.writestr("beach_day/sunset_over_the_sea.jpg", b"")
        zf.writestr("itinerary.txt", "Ubud rice terraces, snorkelling at Nusa Penida, beach clubs in Canggu.")
    with tarfile.open(FILES / "tax_records.tar.gz", "w:gz") as tf:
        data = b"Income tax return, receipts and invoices for the accountant."
        info = tarfile.TarInfo("tax/summary.txt")
        info.size = len(data)
        tf.addfile(info, io.BytesIO(data))
    with py7zr.SevenZipFile(FILES / "wedding.7z", "w") as z:
        z.writestr(b"", "wedding_photos/first_dance.jpg")
        z.writestr(b"", "wedding_photos/bride_and_groom_cake.jpg")
    return 8


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    HOME.mkdir(parents=True, exist_ok=True)
    expected = make_samples()
    step("made sample files", True, expected)
    # Watch only the samples, before the app reads its settings.
    (HOME / "prefs.json").write_text(json.dumps({"folders": [str(FILES)], "name": "Tester"}))

    from PySide6.QtWidgets import QApplication

    import lookhound.app as A
    from lookhound.hunt import Hunter

    app = QApplication([])
    A.load_font(app)
    window = A.Window()
    window.show()
    step("window opened", True)

    def pump(seconds):
        end = time.time() + seconds
        while time.time() < end:
            app.processEvents()
            time.sleep(0.05)

    started = time.time()
    while time.time() - started < INDEX_TIMEOUT:
        pump(1)
        if window.status.text().startswith("All tucked in") and window.shelf.count() >= expected:
            break
    took = round(time.time() - started)
    step("indexed every file", window.shelf.count() >= expected,
         f"{window.shelf.count()} files, status '{window.status.text()}', {took}s")
    window.grab().save(str(OUT / "home.png"))

    hunter = Hunter(window.brain)
    searches = {
        "holiday in Portugal": "trip_notes.md",
        "how to cook spaghetti": "pasta_recipe.txt",
        "a function that checks a password": "login.py",
        "a blue circle": "blue_circle.png",
        "sales presentation": "sales_deck.pptx",
        "snorkelling and surfing trip": "bali_2025.zip",
        "receipts for the accountant": "tax_records.tar.gz",
        "wedding cake photos": "wedding.7z",
    }
    for query, wanted in searches.items():
        found, _ = hunter.search(query)
        names = [r["name"] for r in found[:3]]
        step(f"search '{query}'", names[:1] == [wanted], f"top results {names}")

    found, _ = hunter.search("photos", kinds=A.FILTERS["archive"])
    step("Archives tab shows only archives", found and all(r["kind"] == "archive" for r in found),
         [r["name"] for r in found])

    window.home_query.setText("sales presentation")
    window._go()
    pump(4)
    window.grab().save(str(OUT / "results.png"))
    step("results page shown", window.pages.currentIndex() == 1)
    card = window.list.itemWidget(window.list.item(0)) if window.list.count() else None
    if card:
        card.texts[0].setSelection(0, 5)
    step("result text can be highlighted",
         card is not None and card.selected_text() == "sales", card and card.selected_text())

    report["passed"] = all(s["ok"] for s in report["steps"])


try:
    main()
except Exception:
    step("crashed", False, traceback.format_exc())
finally:
    _save()
    os._exit(0)
